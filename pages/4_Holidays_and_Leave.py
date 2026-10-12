"""Screen D: Holidays & Employee Leave Management.
"""

import streamlit as st
import pandas as pd
import openpyxl
import os
from datetime import date, datetime
from src.state import init_session_state
from src.models import Holiday, HolidayPolicy
from src.configuration import INT_TO_WEEKDAY

init_session_state()

st.title("🏖️ Holidays & Approved Leave")
st.markdown("Manage public holidays, assign holiday staffing policies, and register approved employee leaves.")

config = st.session_state.config
employees = st.session_state.employees

tab1, tab2, tab3 = st.tabs(["🎉 Public Holidays", "✈️ Employee Leave Management", "📥 Upload Holidays Excel"])

with tab1:
    st.subheader("Configured Public Holidays")
    if config.holidays:
        h_data = []
        for h in config.holidays:
            h_data.append({
                "Date": h.date.strftime("%Y-%m-%d"),
                "Day": INT_TO_WEEKDAY[h.date.weekday()],
                "Holiday Name": h.name,
                "Policy": h.policy.value,
                "Description": h.description,
            })
        st.dataframe(pd.DataFrame(h_data), use_container_width=True)
    else:
        st.info("No public holidays configured.")

    st.markdown("#### Add Holiday")
    with st.form("add_holiday_form"):
        h_date = st.date_input("Holiday Date:", value=date(2026, 11, 4))
        h_name = st.text_input("Holiday Name:", value="Deepavali")
        h_policy = st.selectbox("Staffing Policy:", ["SUNDAY", "NORMAL", "CUSTOM"], help="SUNDAY = Apply Sunday minimums; NORMAL = Regular weekday minimums; CUSTOM = Custom numbers.")
        h_desc = st.text_input("Description:", value="Public Festival Holiday")

        c1, c2, c3 = st.columns(3)
        with c1:
            cust_s1 = st.number_input("Custom S1 (if CUSTOM)", min_value=0, max_value=50, value=2)
        with c2:
            cust_s2 = st.number_input("Custom S2 (if CUSTOM)", min_value=0, max_value=50, value=6)
        with c3:
            cust_s3 = st.number_input("Custom S3 (if CUSTOM)", min_value=0, max_value=50, value=6)

        if st.form_submit_button("Add Holiday"):
            custom_mins = {"S1": int(cust_s1), "S2": int(cust_s2), "S3": int(cust_s3)} if h_policy == "CUSTOM" else {}
            new_h = Holiday(
                date=h_date,
                name=h_name,
                description=h_desc,
                policy=HolidayPolicy(h_policy),
                custom_shift_mins=custom_mins,
            )
            # Remove existing holiday on same date if present
            config.holidays = [h for h in config.holidays if h.date != h_date]
            config.holidays.append(new_h)
            st.success(f"Added holiday '{h_name}' on {h_date}!")
            st.rerun()

with tab2:
    st.subheader("Employee Approved Leave")
    emp_options = {f"{e.id} - {e.name}": e for e in employees if e.active}
    if emp_options:
        selected_emp_label = st.selectbox("Select Employee:", list(emp_options.keys()))
        selected_emp = emp_options[selected_emp_label]

        st.write(f"**Current Leave Dates for {selected_emp.id} ({selected_emp.name}):**")
        if selected_emp.approved_leave:
            leave_list = sorted(list(selected_emp.approved_leave))
            st.write(", ".join(d.strftime("%Y-%m-%d") for d in leave_list))
        else:
            st.write("No approved leaves registered.")

        c_l1, c_l2 = st.columns(2)
        with c_l1:
            new_leave_date = st.date_input("Add Leave Date:", value=st.session_state.start_date)
            if st.button("Add Approved Leave"):
                selected_emp.approved_leave.add(new_leave_date)
                st.session_state.current_schedule = None
                st.success(f"Added leave on {new_leave_date} for {selected_emp.id}!")
                st.rerun()
        with c_l2:
            if selected_emp.approved_leave:
                remove_leave_date = st.selectbox("Select Leave Date to Remove:", sorted(list(selected_emp.approved_leave)))
                if st.button("Remove Leave Date"):
                    selected_emp.approved_leave.remove(remove_leave_date)
                    st.session_state.current_schedule = None
                    st.success(f"Removed leave date {remove_leave_date}!")
                    st.rerun()

with tab3:
    st.subheader("Upload Holidays Workbook")
    hol_upload = st.file_uploader("Upload Holidays Excel (.xlsx)", type=["xlsx"])
    if hol_upload is not None:
        try:
            wb = openpyxl.load_workbook(hol_upload, data_only=True)
            ws = wb.active
            imported_count = 0
            for r in range(2, ws.max_row + 1):
                raw_d = ws.cell(row=r, column=1).value
                raw_name = ws.cell(row=r, column=2).value
                raw_desc = ws.cell(row=r, column=3).value
                raw_pol = ws.cell(row=r, column=4).value

                if raw_d is not None and raw_name is not None:
                    if isinstance(raw_d, datetime):
                        dt = raw_d.date()
                    elif isinstance(raw_d, date):
                        dt = raw_d
                    else:
                        dt = datetime.strptime(str(raw_d).strip(), "%Y-%m-%d").date()

                    pol_str = str(raw_pol).strip().upper() if raw_pol else "NORMAL"
                    pol = HolidayPolicy(pol_str) if pol_str in HolidayPolicy.__members__ else HolidayPolicy.NORMAL

                    h_item = Holiday(date=dt, name=str(raw_name), description=str(raw_desc or ""), policy=pol)
                    config.holidays = [h for h in config.holidays if h.date != dt]
                    config.holidays.append(h_item)
                    imported_count += 1
            st.success(f"Imported {imported_count} holidays successfully!")
            st.rerun()
        except Exception as ex:
            st.error(f"Error parsing holidays workbook: {ex}")

    sample_hol_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), "data", "sample_holidays.xlsx")
    if os.path.exists(sample_hol_path):
        with open(sample_hol_path, "rb") as f:
            st.download_button(
                "📥 Download Sample Holidays Excel Template",
                data=f.read(),
                file_name="sample_holidays.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
