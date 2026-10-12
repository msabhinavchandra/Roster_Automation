"""Tests for Excel importing, exporting, multi-size workbooks, and sheet preservation.
"""

import os
import openpyxl
from datetime import date
from src.excel_handler import (
    load_employees_from_excel,
    create_sample_employees_workbook,
    export_roster_to_excel,
)
from src.models import RosterSchedule, SchedulingMode


def test_import_sample_50_employees(tmp_path):
    """Test 1: Importing sample employee workbook with 50 employees."""
    sample_path = str(tmp_path / "sample_50.xlsx")
    create_sample_employees_workbook(sample_path, num_employees=50)

    employees, msgs = load_employees_from_excel(sample_path)
    assert len(employees) == 50
    assert employees[0].id == "EMP001"
    assert employees[-1].id == "EMP050"
    assert all(e.active for e in employees)
    # Check approved leave on EMP005
    emp5 = next(e for e in employees if e.id == "EMP005")
    assert date(2026, 11, 15) in emp5.approved_leave


def test_import_different_sizes(tmp_path):
    """Test 2: Importing employee workbooks of different pool sizes (25 and 65)."""
    for size in [25, 65]:
        p = str(tmp_path / f"emps_{size}.xlsx")
        create_sample_employees_workbook(p, num_employees=size)
        emps, _ = load_employees_from_excel(p)
        assert len(emps) == size


def test_excel_export_and_reopen(tmp_path, sample_50_employees, default_config):
    """Test 20: Generating publication Excel workbook and reopening with openpyxl."""
    schedule = RosterSchedule(
        start_date=date(2026, 11, 1),
        end_date=date(2026, 11, 7),
        employees=sample_50_employees,
        mode=SchedulingMode.MODE_B,
    )
    # Put dummy assignments
    for e in sample_50_employees:
        schedule.set_assignment(e.id, date(2026, 11, 1), "S1")

    out_file = str(tmp_path / "exported_roster.xlsx")
    export_roster_to_excel(
        schedule=schedule,
        config=default_config,
        employees=sample_50_employees,
        output_path=out_file,
    )

    assert os.path.exists(out_file)

    # Reopen and verify required sheets
    wb = openpyxl.load_workbook(out_file)
    expected_sheets = [
        "Final Roster",
        "Employee Summary",
        "Daily Shift Coverage",
        "Night Shift Summary",
        "Validation Report",
        "Configuration Summary",
        "Applied Holidays",
        "Feasibility Report",
    ]
    for es in expected_sheets:
        assert es in wb.sheetnames

    # Check content in Final Roster
    ws = wb["Final Roster"]
    assert ws.cell(row=1, column=1).value == "Employee ID"
    assert ws.cell(row=2, column=1).value == "EMP001"
    assert ws.cell(row=2, column=3).value == "S1"


def test_preservation_of_unrelated_worksheets(tmp_path, sample_50_employees, default_config):
    """Test 21: Preserving unrelated worksheets in an existing workbook."""
    base_file = str(tmp_path / "existing_company_book.xlsx")
    wb = openpyxl.Workbook()
    ws_custom = wb.active
    ws_custom.title = "Department_Budgets"
    ws_custom["A1"] = "Budget 2026"
    ws_custom["B1"] = 500000
    wb.save(base_file)

    schedule = RosterSchedule(
        start_date=date(2026, 11, 1),
        end_date=date(2026, 11, 7),
        employees=sample_50_employees,
    )

    out_file = str(tmp_path / "updated_roster.xlsx")
    export_roster_to_excel(
        schedule=schedule,
        config=default_config,
        employees=sample_50_employees,
        output_path=out_file,
        base_workbook_source=base_file,
    )

    # Reopen and check Department_Budgets is preserved!
    wb_updated = openpyxl.load_workbook(out_file)
    assert "Department_Budgets" in wb_updated.sheetnames
    assert wb_updated["Department_Budgets"]["A1"].value == "Budget 2026"
    assert wb_updated["Department_Budgets"]["B1"].value == 500000
    assert "Final Roster" in wb_updated.sheetnames
